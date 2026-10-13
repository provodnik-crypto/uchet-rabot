"""Проверяет файл отчёта: открывается openpyxl, 5 листов, есть нативные диаграммы, итоги сходятся до копейки."""
import sys
import zipfile
from decimal import Decimal

import openpyxl

path = sys.argv[1] if len(sys.argv) > 1 else 'sample.xlsx'
expected_total = Decimal(sys.argv[2]) if len(sys.argv) > 2 else None

with zipfile.ZipFile(path) as z:
    assert z.testzip() is None
    charts = [n for n in z.namelist() if n.startswith('xl/charts/')]
    images = [n for n in z.namelist() if n.startswith('xl/media/')]
print('диаграмм:', len(charts), '| картинок:', len(images))
assert len(charts) >= 5 and not images, 'диаграммы должны быть объектами Excel, а не картинками'

wb = openpyxl.load_workbook(path)
assert wb.sheetnames == ['Сотрудники', 'Итого за месяц', 'Основной подсчет', 'По видам работ', 'Время'], wb.sheetnames
tv = list(wb['Время'].iter_rows(values_only=True))
assert tv[3][:4] == ('Вид работ', 'Заявок', 'Среднее, ч', 'Медиана, ч'), tv[3]
assert any(r[0] == 'Мастер' for r in tv) and any(r[0] == 'Дольше плана' for r in tv), 'лист «Время»: мастера и заявки дольше плана'
print('лист «Время»:', tv[4])
ws = wb['Основной подсчет']
assert [c.value for c in ws[1]][:10] == ['Дата', 'Отдел', 'Заявка', 'Наименование работ', 'Ед. измерения', 'Кол-во',
                                         'Стоимость за ед.', 'Стоимость общая', 'Примечание', 'Исполнитель']
assert ws.freeze_panes == 'A2' and ws.auto_filter.ref, 'шапка закреплена, автофильтр есть'
rows_total = sum(Decimal(str(r[7])) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None and r[7] is not None and r[0] != 'Итого')
itogo = next(r[7] for r in ws.iter_rows(values_only=True) if r[0] == 'Итого')
s1 = next(r[1] for r in wb['Сотрудники'].iter_rows(values_only=True) if r[0] == 'Итого выполнено работ')
m = next(r[1] for r in wb['Итого за месяц'].iter_rows(values_only=True) if r[0] == 'Итого')
k = next(r[1] for r in wb['По видам работ'].iter_rows(values_only=True) if r[0] == 'Итого')
print('строки:', rows_total, '| итог листа 3:', itogo, '| сотрудники:', s1, '| месяцы:', m, '| виды работ:', k)
for v in (itogo, s1, m, k):
    assert Decimal(str(v)) == rows_total, 'итоги листов не сходятся'
if expected_total is not None:
    assert rows_total == expected_total, f'итог {rows_total} ≠ ожидаемому {expected_total}'
print('OK')
